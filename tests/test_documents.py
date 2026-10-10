"""Tests for binary document text extraction and attachment decoding."""

import io
import struct
from pathlib import Path

import pytest

from mimama.core.context import (
    AttachedFile,
    FileContextError,
    decode_attachment,
    read_files,
)
from mimama.core.documents import (
    DocumentError,
    _clean_doc_text,
    _parse_doc_text,
    extract_text,
    is_document,
)


def _pptx_bytes(text: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _pptx_with_table() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(4), Inches(2)
    ).table
    table.cell(0, 0).text = "H1"
    table.cell(0, 1).text = "H2"
    table.cell(1, 0).text = "R1C1"
    table.cell(1, 1).text = "R1C2"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _docx_bytes(paragraph: str, cells: tuple[str, str] | None = None) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(paragraph)
    if cells is not None:
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = cells[0]
        table.cell(0, 1).text = cells[1]
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Role"])
    ws.append(["Alice", "Dev"])
    ws.append([None, ""])  # exercises the None/empty-cell skip path
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n".encode()
    )
    return bytes(out)


def _synthetic_doc_word(text: bytes) -> bytes:
    """A WordDocument stream whose FIB piece table references ``text`` (ANSI)."""
    csw, cslw = 14, 22
    fib_fc_lcb = 0x22 + csw * 2 + 2 + cslw * 4 + 2
    clx_off = fib_fc_lcb + 268 + 4
    plc_pcd = (
        struct.pack("<II", 0, len(text))
        + struct.pack("<I", (1 << 30) | (clx_off + 1 + 4 + 8 + 8))
        + b"\x00" * 4
    )
    clx = b"\x02" + struct.pack("<I", len(plc_pcd)) + plc_pcd
    return _doc_word_with_clx(clx) + text


def _doc_word_with_clx(clx: bytes) -> bytes:
    """A WordDocument stream whose FIB points at the given CLX piece table."""
    csw, cslw = 14, 22
    fib_fc_lcb = 0x22 + csw * 2 + 2 + cslw * 4 + 2
    clx_off = fib_fc_lcb + 268 + 4
    data = bytearray(fib_fc_lcb + 268 + 4)
    data[0:2] = b"\xec\xa5"
    data[2:4] = struct.pack("<H", 0x00C1)
    data[0x20:0x22] = struct.pack("<H", csw)
    data[0x22 + csw * 2 : 0x22 + csw * 2 + 2] = struct.pack("<H", cslw)
    data[fib_fc_lcb + 264 : fib_fc_lcb + 268] = struct.pack("<I", clx_off)
    data[fib_fc_lcb + 268 : fib_fc_lcb + 272] = struct.pack("<I", len(clx))
    data.extend(clx)
    return bytes(data)


class TestIsDocument:
    @pytest.mark.parametrize("name", ["a.pptx", "a.docx", "a.pdf", "a.xlsx", "a.doc"])
    def test_supported(self, name: str) -> None:
        assert is_document(name)

    @pytest.mark.parametrize("name", ["a.txt", "a.md", "a.png", "a"])
    def test_unsupported(self, name: str) -> None:
        assert not is_document(name)

    def test_extension_is_case_insensitive(self) -> None:
        assert is_document("a.DOCX")


class TestExtractText:
    def test_pptx(self) -> None:
        assert "Slide one body" in extract_text("deck.pptx", _pptx_bytes("Slide one body"))

    def test_docx(self) -> None:
        text = extract_text("notes.docx", _docx_bytes("Para one", ("A", "B")))
        assert "Para one" in text
        assert "A\tB" in text

    def test_pdf(self) -> None:
        assert "Hello PDF" in extract_text("doc.pdf", _minimal_pdf("Hello PDF"))

    def test_xlsx(self) -> None:
        text = extract_text("sheet.xlsx", _xlsx_bytes())
        assert "[Data]" in text
        assert "Name\tRole" in text
        assert "Alice\tDev" in text

    def test_legacy_doc(self) -> None:
        stream = _synthetic_doc_word(b"Hello legacy word")
        assert _parse_doc_text(stream, stream) == "Hello legacy word"

    def test_unsupported_type_raises(self) -> None:
        with pytest.raises(DocumentError, match="unsupported document type"):
            extract_text("notes.txt", b"just text")

    def test_corrupt_pdf_raises_document_error(self) -> None:
        with pytest.raises(DocumentError, match="could not read PDF"):
            extract_text("broken.pdf", b"%PDF-\x00\xffgarbage")

    def test_pptx_extracts_table_cells(self) -> None:
        text = extract_text("table.pptx", _pptx_with_table())
        assert "H1\tH2" in text
        assert "R1C1\tR1C2" in text

    @pytest.mark.parametrize(
        "name,datapath",
        [
            ("bad.pptx", "could not read PowerPoint"),
            ("bad.docx", "could not read Word"),
            ("bad.xlsx", "could not read Excel"),
            ("bad.doc", "could not read legacy .doc"),
        ],
    )
    def test_corrupt_document_raises(self, name: str, datapath: str) -> None:
        with pytest.raises(DocumentError, match=datapath):
            extract_text(name, b"not a real document")


class TestDocParser:
    def test_bad_fib_header(self) -> None:
        with pytest.raises(DocumentError, match="unrecognized .doc FIB header"):
            _parse_doc_text(b"\x00" * 0x30, b"\x00" * 0x30)

    def test_piece_table_not_found(self) -> None:
        stream = _doc_word_with_clx(b"\x03")
        with pytest.raises(DocumentError, match="piece table not found"):
            _parse_doc_text(stream, stream)

    def test_prc_is_skipped(self) -> None:
        # A Prc record (clxt=0x01, empty GrpPrl) that precedes a non-Pcdt
        # yields "piece table not found" (proving the Prc was skipped).
        stream = _doc_word_with_clx(b"\x01\x00\x00\x03")
        with pytest.raises(DocumentError, match="piece table not found"):
            _parse_doc_text(stream, stream)


class TestCleanDocText:
    def test_paragraph_and_cell_marks(self) -> None:
        assert _clean_doc_text("a\rb") == "a\nb"
        assert _clean_doc_text("a\x07b") == "a\tb"

    def test_field_markers_dropped(self) -> None:
        assert _clean_doc_text("a\x13b\x14c\x15d") == "abcd"

    def test_control_characters_dropped(self) -> None:
        assert _clean_doc_text("a\x01\x05b") == "ab"


class TestDecodeAttachment:
    def test_text_passthrough(self) -> None:
        assert decode_attachment("note.txt", b"HELLO") == AttachedFile("note.txt", "HELLO")

    def test_extracts_binary_document(self) -> None:
        attached = decode_attachment("notes.docx", _docx_bytes("Para one"))
        assert attached.name == "notes.docx"
        assert "Para one" in attached.content

    def test_unknown_binary_raises(self) -> None:
        with pytest.raises(FileContextError, match="not a text file"):
            decode_attachment("img.png", b"\x89PNG\r\n\x1a\n\x00\xff")

    def test_size_cap(self) -> None:
        with pytest.raises(FileContextError, match="too large"):
            decode_attachment("big.txt", b"a" * 11, max_chars=10)

    def test_size_cap_applies_to_extracted_text(self) -> None:
        with pytest.raises(FileContextError, match="too large"):
            decode_attachment("notes.docx", _docx_bytes("Para one"), max_chars=2)

    def test_no_extractable_text_raises(self) -> None:
        with pytest.raises(FileContextError, match="no extractable text"):
            decode_attachment("empty.docx", _docx_bytes(""))


class TestReadFiles:
    def test_extracts_docx(self, tmp_path: Path) -> None:
        path = tmp_path / "notes.docx"
        path.write_bytes(_docx_bytes("Server side doc"))
        attached = read_files([str(path)])
        assert attached == [AttachedFile("notes.docx", "Server side doc")]

    def test_rejects_oversized(self, tmp_path: Path) -> None:
        path = tmp_path / "big.txt"
        path.write_text("x" * 11, encoding="utf-8")
        with pytest.raises(FileContextError, match="too large"):
            read_files([str(path)], max_chars=10)
